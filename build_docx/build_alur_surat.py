# -*- coding: utf-8 -*-
"""
Build dokumen "DOKUMENTASI ALUR PELAYANAN SURAT" (SIPADES Desa Karduluk).

Membaca screenshots/manifest.json (dihasilkan screenshot.js) lalu menyusun
dokumen .docx berisi seluruh screenshot alur pengajuan surat per jenis surat
beserta penjelasan pada setiap gambar.

Format mengikuti gaya dokumen skripsi: A4, margin 4/3/3/3 cm,
Times New Roman 12pt, spasi 1,5, caption gambar bernomor, Daftar Gambar
otomatis dengan dot leader + PAGEREF.

Jalankan: python build_docx/build_alur_surat.py
"""
import json
import os
import sys
from PIL import Image
from docx import Document
from docx.shared import Twips, Pt, Emu, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(BASE, ".."))
IMG_DIR = os.path.join(ROOT, "screenshots")
MANIFEST = os.path.join(IMG_DIR, "manifest.json")
LOGO = os.path.join(ROOT, "public", "assets", "logo-karduluk.png")
OUT = os.path.join(IMG_DIR, "DOKUMENTASI-ALUR-SURAT-SIPADES.docx")

# ---------------- state ----------------
IMG_COUNTER = {"n": 0}
FIG_ENTRIES = []
BM_ID = {"n": 1000}
DAFTAR_ELEMS = []

# ---------------- low-level helpers ----------------

def _child(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def set_spacing(p, before=None, after=0, line=360, rule="auto"):
    pPr = p._p.get_or_add_pPr()
    sp = _child(pPr, "w:spacing")
    if before is not None:
        sp.set(qn("w:before"), str(before))
    if after is not None:
        sp.set(qn("w:after"), str(after))
    if line is not None:
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), rule)


def set_indent(p, first_line=None, left=None, hanging=None):
    pPr = p._p.get_or_add_pPr()
    ind = _child(pPr, "w:ind")
    if first_line is not None:
        ind.set(qn("w:firstLine"), str(first_line))
    if left is not None:
        ind.set(qn("w:left"), str(left))
    if hanging is not None:
        ind.set(qn("w:hanging"), str(hanging))


def style_run(r, pt=12, bold=False, italic=False, color=None, name="Times New Roman"):
    r.font.name = name
    rpr = r._r.get_or_add_rPr()
    rf = _child(rpr, "w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width = Twips(11907)
    sec.page_height = Twips(16840)
    sec.top_margin = Twips(1701)
    sec.right_margin = Twips(1701)
    sec.bottom_margin = Twips(1701)
    sec.left_margin = Twips(2268)
    sec.footer_distance = Twips(720)
    sec.header_distance = Twips(720)


def add_page_footer(doc):
    sec = doc.sections[0]
    footer = sec.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        p._element.getparent().remove(p._element)

    sdt = OxmlElement("w:sdt")
    sdtPr = OxmlElement("w:sdtPr")
    sdtId = OxmlElement("w:id"); sdtId.set(qn("w:val"), "-264543138")
    sdtPr.append(sdtId)
    dpo = OxmlElement("w:docPartObj")
    dpg = OxmlElement("w:docPartGallery"); dpg.set(qn("w:val"), "Page Numbers (Bottom of Page)")
    dpu = OxmlElement("w:docPartUnique")
    dpo.append(dpg); dpo.append(dpu); sdtPr.append(dpo)
    sdt.append(sdtPr)
    sdtEnd = OxmlElement("w:sdtEndPr")
    rpr_end = OxmlElement("w:rPr"); rpr_end.append(OxmlElement("w:noProof"))
    sdtEnd.append(rpr_end); sdt.append(sdtEnd)
    sdtContent = OxmlElement("w:sdtContent")

    p_el = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pStyle = OxmlElement("w:pStyle"); pStyle.set(qn("w:val"), "Footer"); pPr.append(pStyle)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); pPr.append(jc)
    p_el.append(pPr)

    def fld_run(ftype):
        r = OxmlElement("w:r"); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), ftype)
        r.append(fc); return r

    p_el.append(fld_run("begin"))
    r_instr = OxmlElement("w:r"); it = OxmlElement("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = " PAGE   \\* MERGEFORMAT "; r_instr.append(it); p_el.append(r_instr)
    p_el.append(fld_run("separate"))
    r_val = OxmlElement("w:r"); rpr_v = OxmlElement("w:rPr"); rpr_v.append(OxmlElement("w:noProof"))
    r_val.append(rpr_v)
    t_v = OxmlElement("w:t"); t_v.text = "1"; r_val.append(t_v); p_el.append(r_val)
    p_el.append(fld_run("end"))
    sdtContent.append(p_el); sdt.append(sdtContent)
    footer._element.append(sdt)


# ---------------- paragraph factories ----------------

def add_title(doc, text, pt=14, before=0, after=120):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    set_spacing(p, before=before, after=after, line=360)
    p.alignment = AL.CENTER
    r = p.add_run(text); style_run(r, pt=pt, bold=True)
    return p


def add_h2(doc, num, title):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    pPr = p._p.get_or_add_pPr()
    sp = _child(pPr, "w:spacing")
    sp.set(qn("w:before"), "240"); sp.set(qn("w:after"), "60")
    sp.set(qn("w:line"), "360"); sp.set(qn("w:lineRule"), "auto")
    r = p.add_run(num + "  " + title); style_run(r, pt=12, bold=True)
    return p


def add_h3(doc, num, title):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 3"]
    pPr = p._p.get_or_add_pPr()
    sp = _child(pPr, "w:spacing")
    sp.set(qn("w:before"), "180"); sp.set(qn("w:after"), "60")
    sp.set(qn("w:line"), "360"); sp.set(qn("w:lineRule"), "auto")
    r = p.add_run(num + "  " + title); style_run(r, pt=12, bold=True)
    return p


def add_body(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    set_spacing(p, before=0, after=0, line=360)
    set_indent(p, first_line=426)
    p.alignment = AL.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        style_run(r_pre, pt=12, bold=True, italic=italic)
    r = p.add_run(text)
    style_run(r, pt=12, italic=italic)
    return p


def add_caption(doc, text, bookmark=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Caption"]
    pPr = p._p.get_or_add_pPr()
    sp = _child(pPr, "w:spacing")
    sp.set(qn("w:before"), "120"); sp.set(qn("w:after"), "120")
    sp.set(qn("w:line"), "360"); sp.set(qn("w:lineRule"), "auto")
    jc = _child(pPr, "w:jc"); jc.set(qn("w:val"), "center")
    ind = _child(pPr, "w:ind"); ind.set(qn("w:firstLine"), "0")

    if bookmark:
        BM_ID["n"] += 1
        st = OxmlElement("w:bookmarkStart")
        st.set(qn("w:id"), str(BM_ID["n"])); st.set(qn("w:name"), bookmark)
        en = OxmlElement("w:bookmarkEnd")
        en.set(qn("w:id"), str(BM_ID["n"]))
        p._p.insert(1, st)
        p._p.append(en)

    if text.startswith("Gambar ") or text.startswith("Tabel "):
        parts = text.split("  ", 1)
        if len(parts) == 2:
            r1 = p.add_run(parts[0] + "  "); style_run(r1, pt=11, bold=True)
            r2 = p.add_run(parts[1]); style_run(r2, pt=11)
            return p

    r = p.add_run(text); style_run(r, pt=11)
    return p


def add_figure(doc, filename, caption_text, explanation):
    IMG_COUNTER["n"] += 1
    num_str = f"Gambar {IMG_COUNTER['n']}"
    full_caption = f"{num_str}  {caption_text}"
    bm = f"gbr_{IMG_COUNTER['n']}"
    FIG_ENTRIES.append((num_str, caption_text, bm))

    path = os.path.join(IMG_DIR, filename)
    try:
        with Image.open(path) as im:
            w_px, h_px = im.size
    except Exception:
        w_px, h_px = 1440, 900

    max_w_emu = int(5.51 * 914400)   # 14 cm
    max_h_emu = int(3.94 * 914400)   # 10 cm
    ratio = w_px / h_px
    w_emu = max_w_emu
    h_emu = int(w_emu / ratio)
    if h_emu > max_h_emu:
        h_emu = max_h_emu
        w_emu = int(h_emu * ratio)

    p = doc.add_paragraph()
    set_spacing(p, before=180, after=60, line=360)
    p.alignment = AL.CENTER
    run = p.add_run()
    run.add_picture(path, width=Emu(w_emu), height=Emu(h_emu))

    add_caption(doc, full_caption, bookmark=bm)

    if explanation:
        ep = doc.add_paragraph()
        ep.style = doc.styles["Normal"]
        set_spacing(ep, before=60, after=0, line=360)
        set_indent(ep, first_line=426)
        ep.alignment = AL.JUSTIFY
        er = ep.add_run(explanation)
        style_run(er, pt=12)

    return p


def add_table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    if widths:
        for i, w in enumerate(widths):
            for cell in tbl.columns[i].cells:
                cell.width = w

    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        set_spacing(p, before=60, after=60, line=240)
        p.alignment = AL.CENTER
        r = p.add_run(h); style_run(r, pt=10, bold=True)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "D9D9D9")
        tcPr.append(shd)

    for ri, row_data in enumerate(rows):
        row = tbl.rows[ri + 1]
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            set_spacing(p, before=40, after=40, line=240)
            p.alignment = AL.LEFT
            r = p.add_run(str(val)); style_run(r, pt=10)

    return tbl


def _page_break_paragraph(doc):
    pb = doc.add_paragraph()
    set_spacing(pb, before=0, after=0, line=240)
    pb.add_run().add_break(WD_BREAK.PAGE)
    return pb


def add_daftar_gambar(doc):
    DAFTAR_ELEMS.append(_page_break_paragraph(doc)._p)

    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    pPr = p._p.get_or_add_pPr()
    sp = _child(pPr, "w:spacing")
    sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "240")
    sp.set(qn("w:line"), "360"); sp.set(qn("w:lineRule"), "auto")
    jc = _child(pPr, "w:jc"); jc.set(qn("w:val"), "center")
    r = p.add_run("DAFTAR GAMBAR"); style_run(r, pt=12, bold=True)
    DAFTAR_ELEMS.append(p._p)

    blank = doc.add_paragraph()
    set_spacing(blank, before=0, after=0, line=360)
    DAFTAR_ELEMS.append(blank._p)

    for num_str, caption_text, bm in FIG_ENTRIES:
        ep = doc.add_paragraph()
        ep.style = doc.styles["Normal"]
        set_spacing(ep, before=0, after=0, line=360)
        epPr = ep._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), "7938")
        tabs.append(tab)
        epPr.append(tabs)

        r1 = ep.add_run(f"{num_str}  {caption_text}")
        style_run(r1, pt=12)
        ep.add_run()._r.append(OxmlElement("w:tab"))

        fld_begin = OxmlElement("w:r"); fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "begin"); fld_begin.append(fc)
        r_instr = OxmlElement("w:r"); it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" PAGEREF {bm} \\h "
        r_instr.append(it)
        r_sep = OxmlElement("w:r"); fc2 = OxmlElement("w:fldChar")
        fc2.set(qn("w:fldCharType"), "separate"); r_sep.append(fc2)
        r_val = ep.add_run("0"); style_run(r_val, pt=12)
        r_end = OxmlElement("w:r"); fc3 = OxmlElement("w:fldChar")
        fc3.set(qn("w:fldCharType"), "end"); r_end.append(fc3)
        for el in (fld_begin, r_instr, r_sep, r_val._r, r_end):
            ep._p.append(el)
        DAFTAR_ELEMS.append(ep._p)

    DAFTAR_ELEMS.append(_page_break_paragraph(doc)._p)


def move_daftar_gambar_after_cover(doc, cover_count):
    body = doc.element.body
    for el in DAFTAR_ELEMS:
        body.remove(el)
    for i, el in enumerate(DAFTAR_ELEMS):
        body.insert(cover_count + i, el)


def enable_update_fields(out_path):
    import zipfile, shutil
    tmp = out_path + ".tmp"
    with zipfile.ZipFile(out_path) as zin:
        infos = zin.infolist()
        contents = {i.filename: zin.read(i.filename) for i in infos}
    settings = contents["word/settings.xml"].decode("utf-8")
    if "updateFields" not in settings:
        settings = settings.replace("</w:settings>", '<w:updateFields w:val="true"/></w:settings>')
        contents["word/settings.xml"] = settings.encode("utf-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in infos:
            zout.writestr(info, contents[info.filename])
    shutil.move(tmp, out_path)


# ---------------- cover ----------------

def add_cover(doc, manifest):
    from datetime import datetime

    if os.path.exists(LOGO):
        p = doc.add_paragraph()
        set_spacing(p, before=1200, after=240, line=240)
        p.alignment = AL.CENTER
        p.add_run().add_picture(LOGO, width=Emu(int(1.38 * 914400)))  # ~3.5 cm

    add_title(doc, "DOKUMENTASI ALUR PELAYANAN SURAT", pt=16, before=240, after=120)
    add_title(doc, "Sistem Informasi Pelayanan Desa (SIPADES)", pt=13, before=0, after=60)
    add_title(doc, "Desa Karduluk, Kecamatan Pragaan, Kabupaten Sumenep", pt=12, before=0, after=480)

    total_fig = len(manifest.get("shared", [])) + sum(
        len(l.get("screenshots", [])) for l in manifest.get("letters", [])
    )
    add_body(doc, f"Dokumen ini memuat dokumentasi visual seluruh alur pelayanan persuratan pada Sistem Informasi Pelayanan Desa (SIPADES) Desa Karduluk, mulai dari pengajuan oleh warga, proses verifikasi dan persetujuan berjenjang oleh perangkat desa (Petugas Desa, Sekretaris Desa, dan Kepala Desa), hingga penerbitan surat resmi beserta verifikasi Tanda Tangan Elektronik (TTE).")
    add_body(doc, f"Dokumen dihasilkan secara otomatis dari {total_fig} tangkapan layar aplikasi pada {manifest.get('generated_at', '')[:10]} dan mencakup {len(manifest.get('letters', []))} jenis surat pelayanan desa.")


def add_pendahuluan(doc, manifest):
    add_h2(doc, "1.", "Pendahuluan")
    add_body(doc, "Pelayanan persuratan SIPADES melibatkan empat peran pengguna. Warga berperan sebagai pemohon melalui portal mandiri, sedangkan tiga peran perangkat desa bertindak sebagai verifikator dan pemberi persetujuan sesuai tingkat kewenangannya. Jenjang persetujuan bersifat dinamis: setiap jenis surat dapat dikonfigurasi dengan 1, 2, atau 3 level approval.")

    add_table(
        doc,
        ["Peran", "Kewenangan dalam Alur Persuratan"],
        [
            ["Warga", "Mendaftar, masuk dengan OTP WhatsApp, mengajukan surat beserta lampiran persyaratan, memantau status, dan mengunduh surat terbit."],
            ["Petugas Desa", "Verifikasi awal (Level 1): memeriksa kelengkapan berkas, menyetujui, meminta revisi, atau menolak permohonan."],
            ["Sekretaris Desa", "Persetujuan administratif (Level 2) untuk jenis surat yang membutuhkan dua atau tiga tingkat persetujuan."],
            ["Kepala Desa", "Persetujuan akhir (Level 3) sekaligus pengesahan dokumen secara elektronik (TTE) dan penerbitan surat."],
        ],
        widths=[Twips(1450), Twips(6400)],
    )
    doc.add_paragraph()

    rows = []
    for l in manifest.get("letters", []):
        rows.append([
            l.get("kode", "-"),
            l.get("nama", "-"),
            f"{l.get('level', '-')} Level",
            "Ya" if l.get("tte") else "Tidak",
            f"{l.get('estimasi', '-')} hari",
            ", ".join(l.get("persyaratan", [])) or "-",
        ])
    add_table(
        doc,
        ["Kode", "Jenis Surat", "Approval", "TTE Kades", "Estimasi", "Persyaratan"],
        rows,
        widths=[Twips(1100), Twips(1500), Twips(650), Twips(620), Twips(600), Twips(3350)],
    )
    doc.add_paragraph()
    add_body(doc, "Catatan: kolom TTE Kades menandakan apakah persetujuan akhir disertai penandatanganan elektronik oleh pejabat berwenang pada level akhir. Pada alur 3 level dokumen ditandatangani Kepala Desa, pada alur 2 level oleh Sekretaris Desa, dan pada alur 1 level oleh Petugas Desa atas nama Kepala Desa. Seluruh dokumen dibekali token verifikasi yang dapat divalidasi melalui halaman publik verifikasi surat.")


def add_shared_section(doc, manifest, section_no):
    add_h2(doc, f"{section_no}.", "Akses Sistem & Autentikasi")
    add_body(doc, "Bagian ini mendokumentasikan halaman pintu masuk layanan serta mekanisme autentikasi yang digunakan warga dan perangkat desa sebelum menjalankan alur persuratan.")

    for s in manifest.get("shared", []):
        add_figure(doc, s["file"], s["caption"], s.get("explanation", ""))


def add_letter_sections(doc, manifest, section_no):
    add_h2(doc, f"{section_no}.", "Alur Pelayanan per Jenis Surat")
    add_body(doc, "Setiap subbagian berikut menampilkan rangkaian tangkapan layar alur lengkap satu jenis surat: pengajuan oleh warga, verifikasi pada setiap level persetujuan, penerbitan surat, hingga verifikasi keabsahan dokumen. Penjelasan disajikan pada setiap gambar.")

    for idx, letter in enumerate(manifest.get("letters", []), start=1):
        if not letter.get("screenshots"):
            continue

        adjectif = f"({letter.get('level', '-')} Level Persetujuan" + (" & TTE)" if letter.get("tte") else ")")
        add_h3(doc, f"{section_no}.{idx}", f"Alur {letter.get('nama', '-')} {adjectif}")
        add_body(doc, f"Jenis surat {letter.get('nama', '-')} (kode {letter.get('kode', '-')}) memiliki {letter.get('level', '-')} tingkat persetujuan dengan estimasi penyelesaian {letter.get('estimasi', '-')} hari kerja. Dokumen persyaratan yang harus dilampirkan: {', '.join(letter.get('persyaratan', [])) or 'tidak ada'}.")

        for s in letter["screenshots"]:
            add_figure(doc, s["file"], s["caption"], s.get("explanation", ""))


def add_penutup(doc, section_no):
    add_h2(doc, f"{section_no}.", "Penutup")
    add_body(doc, "Dokumentasi ini menunjukkan bahwa seluruh alur pelayanan surat pada SIPADES Desa Karduluk berjalan secara konsisten untuk setiap jenis surat: pengajuan mandiri oleh warga melalui portal, verifikasi dan persetujuan berjenjang sesuai konfigurasi level tiap jenis surat, notifikasi WhatsApp pada setiap perubahan status, penerbitan surat resmi berformat PDF, serta verifikasi keabsahan dokumen melalui token TTE pada halaman publik.")
    add_body(doc, "Dokumen dapat diperbarui secara berkala dengan menjalankan ulang generator tangkapan layar (screenshot.js) dan pembangun dokumen ini (build_alur_surat.py) mengikuti panduan pada berkas README terkait.")


# ---------------- main ----------------

def build():
    if not os.path.exists(MANIFEST):
        print(f"Manifest tidak ditemukan: {MANIFEST}")
        print("Jalankan screenshot.js terlebih dahulu untuk menghasilkan manifest.")
        sys.exit(1)

    with open(MANIFEST, "r", encoding="utf-8") as f:
        manifest = json.load(f)

    IMG_COUNTER["n"] = 0
    FIG_ENTRIES.clear()
    BM_ID["n"] = 1000
    DAFTAR_ELEMS.clear()

    doc = Document()
    setup_page(doc)
    add_page_footer(doc)

    add_cover(doc, manifest)
    cover_count = len(doc.paragraphs)  # jumlah paragraf cover (sebelum sisipan Daftar Gambar)

    add_pendahuluan(doc, manifest)
    add_shared_section(doc, manifest, 2)
    add_letter_sections(doc, manifest, 3)
    add_penutup(doc, 4)

    add_daftar_gambar(doc)
    move_daftar_gambar_after_cover(doc, cover_count)

    doc.save(OUT)
    enable_update_fields(OUT)

    total_fig = len(FIG_ENTRIES)
    print(f"✅ Dokumen dibuat: {OUT}")
    print(f"   Total gambar : {total_fig}")
    print(f"   Jenis surat  : {len(manifest.get('letters', []))}")


if __name__ == "__main__":
    build()
